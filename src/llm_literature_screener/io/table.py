"""Reader and writer for CSV, TSV and Excel files with a configurable column mapping.

Columns are matched to record fields by common header names (case-insensitive), for
example ``title``, ``abstract``, ``authors``, ``year``, ``doi`` and ``pmid``. A column
mapping overrides the defaults, for example ``{"title": "Article Title"}``. When records
are written back, every source column is kept in its original order and the tool adds
``lls_score``, ``lls_p_include``, ``lls_decision`` and ``lls_reason``.
"""

from __future__ import annotations

import csv
import datetime as dt
from collections.abc import Callable, Mapping, Sequence
from dataclasses import dataclass, field
from io import StringIO
from pathlib import Path
from typing import TypeAlias

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from llm_literature_screener.io.common import (
    OUTPUT_FIELDS,
    OUTPUT_PREFIX,
    assign_ids,
    clean_text,
    format_label,
    format_number,
    is_output_field,
    normalize_doi,
    normalize_pmid,
    parse_label,
    parse_year,
    read_text,
    split_authors,
)
from llm_literature_screener.record import Record

FIELDS = ("id", "title", "abstract", "authors", "year", "journal", "doi", "pmid", "label")
"""Record fields that can be mapped to a column."""

DEFAULT_COLUMNS: dict[str, tuple[str, ...]] = {
    "id": ("id", "record_id", "record id", "recordid", "ref id", "refid"),
    "title": ("title", "article title", "primary title", "document title", "ti"),
    "abstract": ("abstract", "abstract note", "ab"),
    "authors": ("authors", "author", "author(s)", "author names", "au"),
    "year": (
        "year",
        "publication year",
        "publication_year",
        "published year",
        "pubyear",
        "py",
        "date",
    ),
    "journal": (
        "journal",
        "journal title",
        "source title",
        "publication title",
        "secondary title",
        "source",
        "so",
    ),
    "doi": ("doi", "di"),
    "pmid": ("pmid", "pubmed_id", "pubmed id", "pm"),
    "label": (
        "label",
        "label_included",
        "included",
        "include",
        "decision",
        "screening decision",
        "final decision",
    ),
}
"""Header names tried for each field, in order of preference."""

Cell: TypeAlias = str | float
LabelParser: TypeAlias = Callable[[Mapping[str, str]], bool | None]


class ColumnMappingError(ValueError):
    """Raised when a column mapping names an unknown field or a missing column."""


@dataclass
class Table:
    """A rectangular table of text cells with a header row."""

    headers: list[str]
    rows: list[list[str]] = field(default_factory=list)


@dataclass(frozen=True)
class TableProfile:
    """How the columns of one kind of table map to record fields.

    Attributes:
        name: Short name of the profile, used in error messages.
        columns: Header names tried for each field, in order of preference.
        label_parser: Optional function that reads the human label from a whole row, for
            exports that store decisions in a free-text column (such as Rayyan's notes).
    """

    name: str
    columns: Mapping[str, tuple[str, ...]] = field(default_factory=lambda: DEFAULT_COLUMNS)
    label_parser: LabelParser | None = None


GENERIC_PROFILE = TableProfile(name="table")


def _normalize_header(header: str) -> str:
    return " ".join(header.replace("﻿", "").strip().lower().split())


def _cell_text(value: object) -> str:
    """Convert a spreadsheet cell to text the way a person would read it."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    return str(value)


def _detect_delimiter(header_line: str) -> str:
    counts = {delimiter: header_line.count(delimiter) for delimiter in (",", ";", "\t")}
    best = max(counts, key=lambda delimiter: counts[delimiter])
    return best if counts[best] else ","


def _rectangular(headers: list[str], rows: list[list[str]]) -> Table:
    width = len(headers)
    padded = []
    for row in rows:
        if not any(cell.strip() for cell in row):
            continue
        padded.append((row + [""] * width)[:width])
    return Table(headers=headers, rows=padded)


def read_delimited(path: Path, delimiter: str | None = None) -> Table:
    """Read a CSV or TSV file. The delimiter is detected from the header line when not given."""
    text = read_text(path)
    lines = text.splitlines()
    if lines and lines[0].strip().lower().startswith("sep="):
        delimiter = delimiter or lines[0].strip()[4:5] or ","
        text = text.split("\n", 1)[1] if "\n" in text else ""
    if not text.strip():
        return Table(headers=[])
    if delimiter is None:
        delimiter = _detect_delimiter(text.splitlines()[0])
    reader = csv.reader(StringIO(text, newline=""), delimiter=delimiter)
    rows = [list(row) for row in reader]
    if not rows:
        return Table(headers=[])
    headers = [header.replace("﻿", "").strip() for header in rows[0]]
    return _rectangular(headers, rows[1:])


def read_excel(path: Path, sheet: str | None = None) -> Table:
    """Read the first worksheet of an Excel file, or the worksheet named ``sheet``."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet] if sheet else workbook.worksheets[0]
        rows = [[_cell_text(cell) for cell in row] for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    while rows and not any(cell.strip() for cell in rows[0]):
        rows.pop(0)
    if not rows:
        return Table(headers=[])
    headers = [header.strip() for header in rows[0]]
    while headers and not headers[-1]:
        headers.pop()
    return _rectangular(headers, rows[1:])


def resolve_columns(
    headers: Sequence[str],
    profile: TableProfile = GENERIC_PROFILE,
    mapping: Mapping[str, str] | None = None,
) -> dict[str, int]:
    """Find the column index of each record field.

    Args:
        headers: Header row of the table.
        profile: Default header names for each field.
        mapping: Explicit field-to-header mapping that overrides the defaults.

    Returns:
        A mapping from field name to column index, for the fields that were found.

    Raises:
        ColumnMappingError: If the mapping names an unknown field or a missing column, or if
            neither a title nor an abstract column can be found.
    """
    positions: dict[str, int] = {}
    for index, header in enumerate(headers):
        key = _normalize_header(header)
        if key and not is_output_field(key):
            positions.setdefault(key, index)
    resolved: dict[str, int] = {}
    for field_name, candidates in profile.columns.items():
        for candidate in candidates:
            found = positions.get(_normalize_header(candidate))
            if found is not None:
                resolved[field_name] = found
                break
    for field_name, header in (mapping or {}).items():
        if field_name not in FIELDS:
            raise ColumnMappingError(
                f"Unknown field {field_name!r} in the column mapping; "
                f"use one of: {', '.join(FIELDS)}."
            )
        mapped = positions.get(_normalize_header(header))
        if mapped is None:
            available = ", ".join(repr(h) for h in headers if h)
            raise ColumnMappingError(
                f"Column {header!r} (mapped to {field_name}) is not in the file. "
                f"Available columns: {available}."
            )
        resolved[field_name] = mapped
    if "title" not in resolved and "abstract" not in resolved:
        raise ColumnMappingError(
            f"No title or abstract column found in this {profile.name} file. Map one "
            "explicitly, for example --column title='Article Title'."
        )
    return resolved


def records_from_table(
    table: Table,
    profile: TableProfile = GENERIC_PROFILE,
    mapping: Mapping[str, str] | None = None,
) -> list[Record]:
    """Convert table rows to records using the profile's defaults and an optional mapping."""
    columns = resolve_columns(table.headers, profile, mapping)

    def cell(row: list[str], field_name: str) -> str:
        index = columns.get(field_name)
        return clean_text(row[index]) if index is not None else ""

    ids = assign_ids([cell(row, "id") for row in table.rows])
    records = []
    for row, ident in zip(table.rows, ids, strict=True):
        if profile.label_parser is not None and "label" not in (mapping or {}):
            label = profile.label_parser(dict(zip(table.headers, row, strict=True)))
        else:
            label = parse_label(cell(row, "label")) if "label" in columns else None
        records.append(
            Record(
                id=ident,
                title=cell(row, "title"),
                abstract=cell(row, "abstract"),
                authors=split_authors(cell(row, "authors")),
                year=parse_year(cell(row, "year")),
                journal=cell(row, "journal"),
                doi=normalize_doi(cell(row, "doi")),
                pmid=normalize_pmid(cell(row, "pmid")),
                raw=list(zip(table.headers, row, strict=True)),
                label=label,
            )
        )
    return records


def _plain_fields(record: Record) -> list[tuple[str, str]]:
    """Describe a record that was not read from a table with one column per field."""
    return [
        ("id", record.id),
        ("title", record.title),
        ("abstract", record.abstract),
        ("authors", "; ".join(record.authors)),
        ("year", "" if record.year is None else str(record.year)),
        ("journal", record.journal),
        ("doi", record.doi),
        ("pmid", record.pmid),
        ("label", format_label(record.label)),
    ]


def table_rows(records: Sequence[Record]) -> tuple[list[str], list[list[Cell]]]:
    """Lay out records as a header row and rows, adding the tool's output columns.

    Source columns keep their order; columns written by an earlier run of this tool are
    replaced rather than repeated.
    """
    sources = [
        [(k, v) for k, v in (record.raw or _plain_fields(record)) if not is_output_field(k)]
        for record in records
    ]
    headers: list[str] = [name for name, _ in sources[0]] if sources else []
    for fields in sources[1:]:
        for name, _ in fields:
            if name not in headers:
                headers.append(name)
    assessed = any(record.assessment is not None for record in records)
    output_headers = [OUTPUT_PREFIX + name for name in OUTPUT_FIELDS] if assessed else []
    rows: list[list[Cell]] = []
    for record, fields in zip(records, sources, strict=True):
        if [name for name, _ in fields] == headers:
            row: list[Cell] = [value for _, value in fields]
        else:
            values = dict(reversed(fields))
            row = [values.get(name, "") for name in headers]
        if assessed:
            assessment = record.assessment
            if assessment is None:
                row.extend(["", "", "", ""])
            else:
                row.extend(
                    [
                        assessment.score,
                        assessment.p_include,
                        assessment.decision.value,
                        clean_text(assessment.reason or ""),
                    ]
                )
        rows.append(row)
    return headers + output_headers, rows


def write_delimited(records: Sequence[Record], path: Path, delimiter: str = ",") -> None:
    """Write records to a CSV or TSV file (UTF-8)."""
    headers, rows = table_rows(records)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter=delimiter)
        writer.writerow(headers)
        for row in rows:
            writer.writerow(
                [format_number(cell) if isinstance(cell, float) else cell for cell in row]
            )


def write_excel(records: Sequence[Record], path: Path) -> None:
    """Write records to an Excel workbook with a bold, frozen header row."""
    headers, rows = table_rows(records)
    workbook = Workbook()
    worksheet = workbook.active
    if worksheet is None:  # pragma: no cover - a new workbook always has an active sheet
        worksheet = workbook.create_sheet()
    worksheet.title = "records"
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
    worksheet.freeze_panes = "A2"
    for row in rows:
        worksheet.append(row)
    workbook.save(path)


def read_table_records(
    path: Path,
    fmt: str,
    profile: TableProfile = GENERIC_PROFILE,
    mapping: Mapping[str, str] | None = None,
    sheet: str | None = None,
) -> list[Record]:
    """Read records from a CSV (``fmt="csv"``), TSV (``"tsv"``) or Excel (``"excel"``) file."""
    if fmt == "excel":
        table = read_excel(path, sheet=sheet)
    else:
        table = read_delimited(path, delimiter="\t" if fmt == "tsv" else None)
    if not table.headers:
        return []
    return records_from_table(table, profile, mapping)
