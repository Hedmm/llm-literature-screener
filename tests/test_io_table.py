"""Tests for the CSV, TSV and Excel reader and writer."""

import csv
import datetime as dt
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from llm_literature_screener import Assessment, Decision, Record
from llm_literature_screener.io import (
    ColumnMappingError,
    TableFileError,
    read_records,
    write_records,
)


def test_csv_with_default_column_names(fixtures: Path) -> None:
    records = read_records(fixtures / "sample.csv")
    assert [record.id for record in records] == ["C1", "C2", "C3"]
    first = records[0]
    # "Paper name" is not a known title header, so the title stays empty without a mapping.
    assert first.title == ""
    assert first.abstract == "Synthetic abstract, with a comma."
    assert first.authors == ["Example, Ann", "Sample, Ben"]
    assert first.year == 2021
    assert first.journal == "Journal of Synthetic Examples"
    assert first.doi == "10.5555/synth.0201"
    assert first.pmid == "90000201"
    assert [record.label for record in records] == [True, False, None]
    assert records[1].abstract == "Abstract text that spans two lines."


def test_csv_with_explicit_column_mapping(fixtures: Path) -> None:
    records = read_records(fixtures / "sample.csv", columns={"title": "Paper name"})
    assert records[0].title == "Exercise therapy for chronic low back pain, a synthetic trial"


def test_column_mapping_errors_are_helpful(fixtures: Path) -> None:
    with pytest.raises(ColumnMappingError, match="Unknown field 'summary'"):
        read_records(fixtures / "sample.csv", columns={"summary": "Abstract"})
    with pytest.raises(ColumnMappingError, match="Available columns"):
        read_records(fixtures / "sample.csv", columns={"title": "Headline"})
    with pytest.raises(ColumnMappingError, match="only to CSV"):
        read_records(fixtures / "sample.ris", columns={"title": "TI"})


def test_table_without_title_or_abstract_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "bad.csv"
    path.write_text("name,value\na,b\n", encoding="utf-8")
    with pytest.raises(ColumnMappingError, match="No title or abstract column"):
        read_records(path)


def test_csv_with_semicolons_and_excel_separator_hint(tmp_path: Path) -> None:
    path = tmp_path / "semicolons.csv"
    path.write_text("sep=;\ntitle;abstract;year\nA title;An abstract, with comma;2001\n")
    (record,) = read_records(path)
    assert (record.title, record.abstract, record.year) == (
        "A title",
        "An abstract, with comma",
        2001,
    )


def test_tsv(fixtures: Path) -> None:
    records = read_records(fixtures / "sample.tsv")
    assert [(r.id, r.label) for r in records] == [("T1", True), ("T2", False)]
    assert records[1].pmid == "90000402"


def test_csv_round_trip_keeps_columns_and_appends_output(fixtures: Path, tmp_path: Path) -> None:
    records = read_records(fixtures / "sample.csv")
    for index, record in enumerate(records):
        record.assessment = Assessment(
            score=float(index), decision=Decision.UNCERTAIN, reason="Why."
        )
    out = tmp_path / "out.csv"
    write_records(records, out, "csv")
    with out.open(encoding="utf-8", newline="") as handle:
        rows = list(csv.reader(handle))
    assert rows[0] == [
        "Record ID",
        "Paper name",
        "Abstract",
        "Author(s)",
        "Publication Year",
        "Source title",
        "DOI",
        "PMID",
        "Included",
        "Notes",
        "lls_score",
        "lls_p_include",
        "lls_decision",
        "lls_reason",
    ]
    assert rows[1][9] == "keep this column"
    assert rows[2][2] == "Abstract text that\nspans two lines."
    third = records[2].assessment
    assert third is not None
    assert rows[3][10:] == ["2.0", repr(third.p_include), "uncertain", "Why."]

    # A second run replaces the output columns instead of adding new ones.
    again = read_records(out)
    for record in again:
        record.assessment = Assessment(score=-1.0, decision=Decision.EXCLUDE)
    write_records(again, out, "csv")
    with out.open(encoding="utf-8", newline="") as handle:
        header = next(csv.reader(handle))
    assert header.count("lls_score") == 1


def _make_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Title", "Abstract", "Year", "DOI", "label", "Added"])
    sheet.append(["Synthetic Excel record", "Abstract.", 2020.0, "10.5555/synth.0701", "yes", None])
    sheet.append([None, None, None, None, None, None])
    sheet.append(["Second record", None, 2015, None, 0, dt.datetime(2024, 1, 31)])
    workbook.save(path)


def test_excel_reader_converts_cells_to_text(tmp_path: Path) -> None:
    path = tmp_path / "records.xlsx"
    _make_workbook(path)
    records = read_records(path)
    assert [record.title for record in records] == ["Synthetic Excel record", "Second record"]
    assert [record.year for record in records] == [2020, 2015]
    assert [record.label for record in records] == [True, False]
    assert records[0].doi == "10.5555/synth.0701"
    assert records[1].raw_value("Added") == "2024-01-31"
    assert [record.id for record in records] == ["1", "2"]


def test_excel_round_trip(tmp_path: Path) -> None:
    path = tmp_path / "records.xlsx"
    _make_workbook(path)
    records = read_records(path)
    records[0].assessment = Assessment(score=3.25, decision=Decision.INCLUDE)
    records[1].assessment = Assessment(score=-7.5, decision=Decision.EXCLUDE)
    out = tmp_path / "out.xlsx"
    write_records(records, out, "excel")
    workbook = load_workbook(out)
    sheet = workbook.active
    assert sheet is not None
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][-4:] == ("lls_score", "lls_p_include", "lls_decision", "lls_reason")
    assert rows[1][0] == "Synthetic Excel record"
    assert rows[1][-4] == 3.25
    assert rows[2][-2] == "exclude"
    assert [record.title for record in read_records(out)] == [
        "Synthetic Excel record",
        "Second record",
    ]


def test_records_without_source_fields_get_standard_columns(tmp_path: Path) -> None:
    record = Record(id="N1", title="Built in code", authors=["A", "B"], year=1999, label=True)
    out = tmp_path / "new.tsv"
    write_records([record], out, "tsv")
    (parsed,) = read_records(out)
    assert (parsed.id, parsed.title, parsed.authors, parsed.year, parsed.label) == (
        "N1",
        "Built in code",
        ["A", "B"],
        1999,
        True,
    )


def test_unreadable_spreadsheets_raise_a_clear_error(tmp_path: Path) -> None:
    fake = tmp_path / "fake.xlsx"
    fake.write_text("this is not a workbook")
    with pytest.raises(TableFileError, match="not a readable Excel"):
        read_records(fake)
    real = tmp_path / "records.xlsx"
    _make_workbook(real)
    with pytest.raises(TableFileError, match="no worksheet 'Missing'"):
        read_records(real, sheet="Missing")
